"""Liest die PE-Excel ("PE - ...xlsx") eines Projekts.

Die Spalten werden über ihre Überschriften gefunden (Zeile mit "Adresse:"),
die Projektparameter über ihre Beschriftungen ("Immo-Zins:", "Soll-Miete:" ...).
Dadurch funktioniert der Parser auch, wenn sich Spalten verschieben.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl


def _norm(label) -> str:
    if label is None:
        return ""
    return re.sub(r"[\s:]+", " ", str(label)).strip().lower()


# Spaltenüberschrift (normalisiert) -> Feldname
UNIT_COLUMNS = {
    "adresse": "adresse",
    "einheit": "einheit",
    "we-nr.": "we_nr",
    "status": "status",
    "größe": "flaeche",
    "mea": "mea",
    "aktuelle kaltmiete": "ist_miete",
    "1. erhöhung auf": "erh1_qm",
    "2. erhöhung durch reno": "erh2_qm",
    "3. erhöhung auf": "erh3_qm",
    "vorgegebene soll-miete": "soll_qm_vorgabe",
    "vorgegebene subvention": "subvention_vorgabe",
    "veränderte soll-miete": "soll_qm",
    "laufzeit der verlängerten subvention": "subvention_verl_monate",
    "zusätzliche subvention": "subvention_zusatz",
    "gesamte subvention": "subvention",
    "gesamte reno-rückstellung se": "reno_se",
    "sanierung weg": "sanierung_weg",
    "kaufpreis": "kaufpreis_netto",
    "knk/interne rückstellung": "knk",
    "beurkundeter kaufpreis inkl. knk": "kaufpreis",
    "miet-rendite": "mietrendite",
    "miete": "miete_soll",
    "bankrate": "bankrate",
    "hv": "hv",
    "mv": "mv",
    "sonstige nicht-umlagefähige kosten": "sonstige",
    "rücklagen": "ruecklage",
    "brutto cf": "brutto_cf",
    "abschreibung": "afa_monat",
    "steuererstattung p.m.": "steuer_monat",
    "netto cf": "netto_cf",
}


@dataclass
class Params:
    zins: float
    tilgung: float
    rnd_jahre: float
    grenzsteuer: float
    soll_miete_qm: float
    knk_quote: float
    grundstueck_qm: float
    bodenrichtwert: float
    gesamtflaeche: float
    steuermodell: bool
    reno_se_qm: float | None = None


@dataclass
class Massnahme:
    name: str
    betrag: float | None
    je_qm: float | None = None


@dataclass
class Project:
    titel: str
    params: Params
    units: list[dict]
    weg_massnahmen: list[Massnahme] = field(default_factory=list)
    weg_summe: float = 0.0
    se_massnahmen: list[Massnahme] = field(default_factory=list)
    se_summe: float = 0.0

    def unit(self, key: str) -> dict:
        """Einheit über WE-Nr., Einheit-Nr. oder (Teil der) Adresse finden."""
        k = _norm(key)
        for u in self.units:
            if k in (_norm(_fmt_num(u.get("we_nr"))), _norm(_fmt_num(u.get("einheit")))):
                return u
        hits = [u for u in self.units if k in _norm(u["adresse"])]
        if len(hits) == 1:
            return hits[0]
        exact = [u for u in hits if _norm(u["adresse"]) == k]
        if len(exact) == 1:
            return exact[0]
        raise KeyError(f"Einheit '{key}' nicht eindeutig gefunden ({len(hits)} Treffer)")


def _fmt_num(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def _find_labels(ws, max_row=12):
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                out.setdefault(_norm(c.value), []).append(c)
    return out


def _right(ws, cell, offset=1):
    return ws.cell(row=cell.row, column=cell.column + offset).value


def _block(ws, head_cell) -> list[Massnahme]:
    """Liest Maßnahmen-Liste unterhalb einer Überschrift (Name | Betrag | €/m² | pauschal)."""
    items = []
    r = head_cell.row + 1
    col = head_cell.column
    while True:
        name = ws.cell(row=r, column=col).value
        if not isinstance(name, str) or not name.strip():
            break
        betrag = ws.cell(row=r, column=col + 1).value
        je_qm = ws.cell(row=r, column=col + 2).value
        items.append(Massnahme(name.strip().rstrip(":").strip(), betrag, je_qm))
        r += 1
    return items


def load(path: str) -> Project:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    labels = _find_labels(ws)

    def val(label, idx=0, offset=1):
        cells = labels.get(_norm(label))
        if not cells:
            return None
        return _right(ws, cells[idx], offset)

    steuermodell = str(val("Steuerrückerstattungsmodell?") or "").strip().lower() == "ja"
    # Vollfinanzierung (Spalte F/G) vs. Steuersparmodell (Spalte I/J): je 1./2. Vorkommen
    idx = 1 if steuermodell else 0

    # Header-Zeile der Einheitentabelle suchen
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=40):
        for c in row:
            if _norm(c.value) == "adresse":
                header_row = c.row
                break
        if header_row:
            break
    if not header_row:
        raise ValueError("Einheitentabelle (Spalte 'Adresse:') nicht gefunden")

    colmap = {}
    for c in ws[header_row]:
        key = UNIT_COLUMNS.get(_norm(c.value))
        if key and key not in colmap:
            colmap[key] = c.column

    units = []
    r = header_row + 2  # Einheitenzeile unter der Zeile mit (m²)/(€)
    while True:
        adr = ws.cell(row=r, column=colmap["adresse"]).value
        if not adr:
            break
        u = {k: ws.cell(row=r, column=c).value for k, c in colmap.items()}
        units.append(u)
        r += 1
    gesamtflaeche = ws.cell(row=r, column=colmap["flaeche"]).value or sum(u["flaeche"] for u in units)

    sollm = val("Soll-Miete")
    params = Params(
        zins=val("Immo-Zins", idx),
        tilgung=val("Immo-Tilgung", idx),
        rnd_jahre=val("RND (Jahre)"),
        grenzsteuer=val("Grenzsteuer"),
        soll_miete_qm=sollm,
        knk_quote=val("Verkauf - KNK"),
        grundstueck_qm=val("Grundstücksgröße"),
        bodenrichtwert=val("Bodenrichtwert"),
        gesamtflaeche=gesamtflaeche,
        steuermodell=steuermodell,
    )

    weg_head = labels.get(_norm("Baukosten Gemeinschaftsflächen"))
    se_head = labels.get(_norm("Baukosten Sondereigentum"))
    weg = _block(ws, weg_head[0]) if weg_head else []
    se = _block(ws, se_head[0]) if se_head else []
    weg_summe = _right(ws, weg_head[0]) if weg_head else 0
    se_summe = _right(ws, se_head[0]) if se_head else 0
    if se_head:
        params.reno_se_qm = _right(ws, se_head[0], 2)

    titel = ws["C2"].value or ""
    return Project(titel, params, units, weg, weg_summe or 0, se, se_summe or 0)
